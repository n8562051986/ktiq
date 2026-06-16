"""
KTI Generator Library
"""
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(4)
    s.bottom_margin = Cm(3)
    s.left_margin = Cm(4)
    s.right_margin = Cm(3)

def heading(text, level=1):
    p = doc.add_paragraph()
    r = p.add_run(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if level == 0 else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14 if level == 0 else 12)
    r.font.bold = True

def par(text, bold=False, indent=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.0)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.bold = bold

def arap(arabic, trans):
    """Add Arabic + translation paragraphs"""
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(1.0)
    r = p.add_run(arabic)
    r.font.name = 'Traditional Arabic'
    r.font.size = Pt(16)
    rPr = r._element.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:ascii'), 'Traditional Arabic')
    rf.set(qn('w:hAnsi'), 'Traditional Arabic')
    if trans:
        p2 = doc.add_paragraph()
        p2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p2.paragraph_format.space_after = Pt(6)
        p2.paragraph_format.line_spacing = 1.5
        p2.paragraph_format.first_line_indent = Cm(1.0)
        r2 = p2.add_run(trans)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)

def ref_line(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.first_line_indent = Cm(-1.0)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)

def blank():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5

def save(path):
    doc.save(path)
